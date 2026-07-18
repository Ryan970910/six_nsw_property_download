from __future__ import annotations

import csv
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo


BASE_NAME = "weekly_dat_source_deduped_20260105_20260629"
OUTPUT_DIR = Path("outputs/weekly_dat_new_records")
RECORDS_CSV = OUTPUT_DIR / f"{BASE_NAME}.csv"
SUMMARY_CSV = OUTPUT_DIR / f"{BASE_NAME}.summary.csv"
OUTPUT_XLSX = OUTPUT_DIR / f"{BASE_NAME}.xlsx"


def main() -> None:
    wb = Workbook(write_only=True)
    write_csv_sheet(wb, "Source Deduped Records", RECORDS_CSV)
    write_csv_sheet(wb, "Summary", SUMMARY_CSV)
    write_notes_sheet(wb)
    wb.save(OUTPUT_XLSX)

    style_workbook(OUTPUT_XLSX)
    verify_workbook(OUTPUT_XLSX)
    print(f"Saved {OUTPUT_XLSX}")


def write_csv_sheet(wb: Workbook, title: str, csv_path: Path) -> None:
    ws = wb.create_sheet(title)
    with csv_path.open("r", newline="", encoding="utf-8-sig") as file:
        reader = csv.reader(file)
        for row in reader:
            ws.append(row)


def write_notes_sheet(wb: Workbook) -> None:
    ws = wb.create_sheet("Notes")
    ws.append(["Workbook", "NSW weekly DAT source-deduped records"])
    ws.append(["Date range", "2026-01-05 to 2026-06-29"])
    ws.append(["Duplicate handling", "Removed duplicate rows inside each weekly source batch only."])
    ws.append(["Database upload/check", "No database records were uploaded. Database duplicate check was not run because real DB settings were not available in .env.example."])


def style_workbook(path: Path) -> None:
    wb = load_workbook(path)
    header_fill = PatternFill("solid", fgColor="1F4E79")
    header_font = Font(bold=True, color="FFFFFF")

    for ws in wb.worksheets:
        if ws.max_row >= 1:
            for cell in ws[1]:
                cell.fill = header_fill
                cell.font = header_font
            ws.freeze_panes = "A2"
            ws.sheet_view.showGridLines = False

        widths = {}
        for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, 200)):
            for cell in row:
                value = "" if cell.value is None else str(cell.value)
                widths[cell.column_letter] = min(max(widths.get(cell.column_letter, 0), len(value) + 2), 45)
        for column_letter, width in widths.items():
            ws.column_dimensions[column_letter].width = width

    records = wb["Source Deduped Records"]
    records.auto_filter.ref = records.dimensions
    if records.max_row > 1:
        records.add_table(make_table("WeeklyDatSourceDedupedRecords", records.dimensions))

    summary = wb["Summary"]
    summary.auto_filter.ref = summary.dimensions
    if summary.max_row > 1:
        summary.add_table(make_table("WeeklyDatSourceDedupedSummary", summary.dimensions))

    notes = wb["Notes"]
    notes.auto_filter.ref = notes.dimensions

    wb.save(path)


def make_table(name: str, ref: str) -> Table:
    table = Table(displayName=name, ref=ref)
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    return table


def verify_workbook(path: Path) -> None:
    wb = load_workbook(path, read_only=True, data_only=True)
    records = wb["Source Deduped Records"]
    summary = wb["Summary"]
    notes = wb["Notes"]
    print(f"records_rows={records.max_row - 1}")
    print(f"summary_rows={summary.max_row - 1}")
    print(f"notes_rows={notes.max_row}")
    print(f"first_record={next(records.iter_rows(min_row=2, max_row=2, values_only=True))}")
    wb.close()


if __name__ == "__main__":
    main()
